from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from sigit.analitica.consultas import calcular_pivote, indicadores_clave
from sigit.analitica.exportar import libro_xlsx, respuesta_csv
from sigit.jornadas.models import Jornada, JornadaPoblacion
from sigit.nucleo import models as m

from .conftest import como

pytestmark = pytest.mark.django_db


@pytest.fixture
def datos(crear_jornada):
    grupo = m.GrupoPoblacional.objects.get(codigo="ADULTOS")
    a = crear_jornada("BIM23", fecha=date(2025, 5, 1))
    b = crear_jornada("BIM23", lugar="Escuela rural", fecha=date(2026, 2, 1))
    c = crear_jornada("BIM24", lugar="Muelle de Cartagena", fecha=date(2026, 3, 1), dane="13001")
    for jornada, personas in ((a, 10), (b, 20), (c, 30)):
        JornadaPoblacion.objects.create(jornada=jornada, grupo=grupo, cantidad_personas=personas)


def test_pivote_personas_por_anio_y_unidad(datos):
    tabla = calcular_pivote(Jornada.vigentes.all(), "anio", "unidad", "personas")
    assert tabla.filas == ["2025", "2026"]
    assert tabla.columnas == ["BIM23", "BIM24"]
    assert tabla.celdas[("2026", "BIM23")] == 20
    assert tabla.total_fila["2026"] == 50
    assert tabla.total == 60


def test_los_totales_distintos_no_se_suman(datos):
    """Municipios atendidos: la misma Tumaco en dos años cuenta UNA vez en el total."""
    tabla = calcular_pivote(Jornada.vigentes.all(), "unidad", "anio", "municipios")
    assert tabla.total_fila["BIM23"] == 1
    assert sum(tabla.celdas.get(("BIM23", a), 0) for a in tabla.columnas) == 2
    assert tabla.total == 2


def test_el_pivote_respeta_el_ambito(datos, unidades):
    como(unidades["BIM24"])
    assert calcular_pivote(Jornada.vigentes.all(), "unidad", None, "jornadas").filas == ["BIM24"]


def test_indicadores_clave(datos):
    kpi = indicadores_clave(Jornada.vigentes.all())
    assert kpi["total"] == 3 and kpi["personas"] == 60 and kpi["municipios"] == 2


def test_el_tablero_y_la_tabla_dinamica_se_pintan(datos, ingresar):
    cliente = ingresar("JACID")
    assert "Tablero de gestión" in cliente.get("/analitica/tablero/").content.decode()
    r = cliente.get("/analitica/tablas-dinamicas/?medida=personas&fila=departamento&columna=anio")
    assert r.status_code == 200 and "datos-pivote" in r.content.decode()
    xlsx = cliente.get("/analitica/tablas-dinamicas/exportar.xlsx?medida=personas&fila=unidad")
    hoja = load_workbook(io.BytesIO(xlsx.content)).active
    assert hoja["A1"].value == "Unidad" and hoja.cell(hoja.max_row, 1).value == "Total"


def test_informe_guardado(datos, ingresar):
    cliente = ingresar("BIM23")
    cliente.post(
        "/analitica/tablas-dinamicas/informes/", {"nombre": "Mi informe", "parametros": "fila=anio"}
    )
    assert "Mi informe" in cliente.get("/analitica/tablas-dinamicas/").content.decode()


def test_la_exportacion_escapa_formulas():
    """Un texto digitado como «=HYPERLINK(…)» no se convierte en fórmula en Excel."""
    csv = respuesta_csv("x", ["Lugar"], [['=HYPERLINK("http://x")'], ["+57 300"], ["Tumaco"]])
    texto = csv.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in texto and "'+57 300" in texto and "\nTumaco" in texto.replace("\r", "")
    hoja = load_workbook(io.BytesIO(libro_xlsx("h", ["Lugar"], [["=1+1"]]))).active
    assert hoja["A2"].value == "'=1+1"


def test_exportar_jornadas_filtradas(datos, ingresar):
    r = ingresar("JACID").get("/jornadas/exportar.csv?desde=2026-01-01")
    lineas = r.content.decode("utf-8-sig").strip().splitlines()
    assert len(lineas) == 3  # Encabezado + las dos de 2026.
