"""
Exportación a CSV y XLSX.

CSV con BOM y punto y coma: es lo que Excel en español abre sin asistente.
XLSX con openpyxl (licencia MIT), encabezado con la paleta ARC, filtros
automáticos y panel inmovilizado, para que el consolidado se pueda trabajar
como tabla dinámica en Excel sin prepararlo a mano.

Una celda que empieza por = + - @ se escapa: un texto como «=HYPERLINK(...)»
digitado en un campo libre no debe convertirse en fórmula al abrir el archivo
(inyección de fórmulas en CSV).
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PELIGROSOS = ("=", "+", "-", "@", "\t", "\r")


def _seguro(valor: object) -> object:
    if isinstance(valor, str) and valor.startswith(PELIGROSOS):
        return "'" + valor
    return valor


def _texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, Decimal):
        return format(valor, "f")
    return str(_seguro(valor))


def respuesta_csv(
    nombre: str, encabezados: Sequence[str], filas: Iterable[Sequence[object]]
) -> HttpResponse:
    salida = io.StringIO()
    salida.write("﻿")
    escritor = csv.writer(salida, delimiter=";")
    escritor.writerow(encabezados)
    for fila in filas:
        escritor.writerow([_texto(v) for v in fila])
    respuesta = HttpResponse(salida.getvalue(), content_type="text/csv; charset=utf-8")
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}.csv"'
    return respuesta


def libro_xlsx(hoja: str, encabezados: Sequence[str], filas: Iterable[Sequence[object]]) -> bytes:
    libro = Workbook()
    ws = libro.active
    ws.title = hoja[:31]
    ws.append(list(encabezados))
    relleno = PatternFill("solid", fgColor="00205B")
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        ws.append(
            [_seguro(v) if not isinstance(v, datetime) else v.replace(tzinfo=None) for v in fila]
        )
    for indice, titulo in enumerate(encabezados, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = max(12, min(45, len(titulo) + 6))
        for (celda,) in ws.iter_rows(min_row=2, min_col=indice, max_col=indice):
            if isinstance(celda.value, date):
                celda.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def respuesta_xlsx(
    nombre: str, hoja: str, encabezados: Sequence[str], filas: Iterable[Sequence[object]]
) -> HttpResponse:
    respuesta = HttpResponse(
        libro_xlsx(hoja, encabezados, filas),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}.xlsx"'
    return respuesta
